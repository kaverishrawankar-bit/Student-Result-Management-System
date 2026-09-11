import openpyxl
import os

FILE_NAME = "/storage/emulated/0/student_results.xlsx"


# Create Excel file
def create_excel():
    if not os.path.exists(FILE_NAME):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Student Results"

        ws.append([
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ])

        wb.save(FILE_NAME)


# Calculate result
def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    # Check whether student passed all subjects
    if any(mark < 35 for mark in marks):
        grade = "F"
        status = "FAIL"

    elif percentage >= 90:
        grade = "A+"
        status = "PASS"

    elif percentage >= 80:
        grade = "A"
        status = "PASS"

    elif percentage >= 70:
        grade = "B"
        status = "PASS"

    elif percentage >= 60:
        grade = "C"
        status = "PASS"

    elif percentage >= 50:
        grade = "D"
        status = "PASS"

    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


# Add student
def add_student():

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    print("\n----- ADD STUDENT RESULT -----")

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    course = input("Enter Course/Class: ")

    marks = []

    for i in range(1, 6):
        while True:
            try:
                mark = float(input("Enter marks for Subject " + str(i) + ": "))

                if 0 <= mark <= 100:
                    marks.append(mark)
                    break
                else:
                    print("Enter marks between 0 and 100.")

            except ValueError:
                print("Please enter a valid number.")

    total, percentage, grade, status = calculate_result(marks)

    ws.append([
        roll_no,
        name,
        course,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    wb.save(FILE_NAME)

    print("\nStudent result saved successfully!")
    print("Total      :", total)
    print("Percentage :", format(percentage, ".2f") + "%")
    print("Grade      :", grade)
    print("Status     :", status)


# Search student
def get_result():

    roll_no = input("\nEnter Roll No: ")

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    found = False

    for row in ws.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == roll_no:

            print("\n--------------------------------")
            print("         STUDENT RESULT")
            print("--------------------------------")
            print("Roll No    :", row[0])
            print("Name       :", row[1])
            print("Class      :", row[2])
            print("Total      :", row[8])
            print("Percentage :", format(row[9], ".2f") + "%")
            print("Grade      :", row[10])
            print("Status     :", row[11])
            print("--------------------------------")

            found = True
            break

    if not found:
        print("\nStudent record not found.")


# Show all students
def show_all_data():

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    print("\n========== ALL STUDENT DATA ==========")

    found = False

    for row in ws.iter_rows(min_row=2, values_only=True):

        found = True

        print("\nRoll No    :", row[0])
        print("Name       :", row[1])
        print("Class      :", row[2])
        print("Total      :", row[8])
        print("Percentage :", format(row[9], ".2f") + "%")
        print("Grade      :", row[10])
        print("Status     :", row[11])
        print("-------------------------------------")

    if not found:
        print("No student records available.")


# Main menu
def menu():

    create_excel()

    while True:

        print("\n")
        print("================================")
        print("   STUDENT RESULT MANAGEMENT")
        print("================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("\nProgram closed.")
            break

        else:
            print("\nInvalid choice. Please try again.")


# Start program
menu()